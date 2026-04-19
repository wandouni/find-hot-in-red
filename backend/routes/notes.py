import io
import json
from typing import Optional, Literal
from urllib.parse import quote
from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from database import get_db
from models import Note, Comment, Task
from schemas import NoteOut, NoteDetail, CommentOut

router = APIRouter()

DATE_FILTER_LABEL = {0: "不限", 1: "1天内", 2: "2天内", 7: "1周内"}


def _scheme_filename(task: Task | None, ext: str) -> str:
    """Build filename: YYYYMMDD_kw1-kw2_period.ext"""
    if task is None:
        return f"xhs_notes.{ext}"
    date_str = task.created_at.strftime("%Y%m%d") if task.created_at else "unknown"
    try:
        kws = json.loads(task.keywords or "[]")
        kw_str = "-".join(kws) if kws else "全部"
    except Exception:
        kw_str = "全部"
    period = DATE_FILTER_LABEL.get(task.date_filter or 0, "不限")
    return f"{date_str}_{kw_str}_{period}.{ext}"


@router.get("/api/notes")
def list_notes(
    task_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    sort: Literal["likes", "collects", "date"] = Query("date"),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
):
    q = db.query(Note)
    if task_id is not None:
        q = q.filter(Note.task_id == task_id)
    if keyword:
        q = q.filter(Note.keyword == keyword)

    sort_col = {
        "likes": desc(Note.likes),
        "collects": desc(Note.collects),
        "date": desc(Note.crawl_time),
    }[sort]
    q = q.order_by(sort_col)

    total = q.count()
    notes = q.offset(offset).limit(limit).all()
    return {"total": total, "items": [NoteOut.model_validate(n) for n in notes]}


@router.get("/api/export/notes.md")
def export_notes_md(
    task_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Export notes as Markdown sorted by likes descending."""
    task = db.get(Task, task_id) if task_id else None

    q = db.query(Note).order_by(desc(Note.likes))
    if task_id is not None:
        q = q.filter(Note.task_id == task_id)
    if keyword:
        q = q.filter(Note.keyword == keyword)
    notes = q.all()

    lines = []
    for note in notes:
        if note.title:
            title = note.title
        else:
            preview = (note.content or "").strip()[:20]
            title = f"[无标题]{preview}"
        url = note.url or ""
        likes = note.likes or 0
        collects = note.collects or 0
        lines.append(f"[{title}]({url})")
        lines.append(f"点赞数量：{likes}；收藏数量：{collects}；8H热度：；")
        lines.append("")

    content = "\n".join(lines)
    filename = _scheme_filename(task, "md")
    return StreamingResponse(
        iter([content]),
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/api/export/notes")
def export_notes(
    task_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Export notes (optionally filtered by scheme/keyword) as Excel."""
    task = db.get(Task, task_id) if task_id else None

    q = db.query(Note).order_by(desc(Note.crawl_time))
    if task_id is not None:
        q = q.filter(Note.task_id == task_id)
    if keyword:
        q = q.filter(Note.keyword == keyword)
    notes = q.all()

    note_ids = [n.id for n in notes]
    comments_by_note: dict[str, list[Comment]] = {}
    if note_ids:
        all_comments = (
            db.query(Comment)
            .filter(Comment.note_id.in_(note_ids))
            .order_by(Comment.note_id, Comment.rank)
            .all()
        )
        for c in all_comments:
            comments_by_note.setdefault(c.note_id, []).append(c)

    wb = openpyxl.Workbook()

    # ── Sheet 1: Notes ───────────────────────────────────────────────
    ws_notes = wb.active
    ws_notes.title = "笔记"

    header_fill = PatternFill("solid", fgColor="C00000")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    note_headers = [
        "笔记ID", "关键词", "标题", "作者", "作者ID",
        "点赞数", "收藏数", "8H", "发布时间", "抓取时间", "链接", "正文",
    ]
    ws_notes.append(note_headers)
    for col, _ in enumerate(note_headers, 1):
        cell = ws_notes.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for note in notes:
        ws_notes.append([
            note.id,
            note.keyword or "",
            note.title or "",
            note.author or "",
            note.author_id or "",
            note.likes,
            note.collects,
            "",          # 8H — blank placeholder
            note.publish_date or "",
            str(note.crawl_time)[:19] if note.crawl_time else "",
            note.url or "",
            note.content or "",
        ])

    col_widths = [26, 12, 40, 14, 24, 8, 8, 8, 14, 18, 50, 60]
    for i, w in enumerate(col_widths, 1):
        ws_notes.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws_notes.row_dimensions[1].height = 20

    ws_notes.freeze_panes = "D2"
    ws_notes.auto_filter.ref = ws_notes.dimensions

    # ── Sheet 2: Comments ────────────────────────────────────────────
    ws_cmts = wb.create_sheet("评论")
    comment_headers = ["笔记ID", "笔记标题", "评论排名", "评论内容", "评论点赞数"]
    ws_cmts.append(comment_headers)
    for col, _ in enumerate(comment_headers, 1):
        cell = ws_cmts.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    note_title_map = {n.id: (n.title or "") for n in notes}
    for note_id, cmts in comments_by_note.items():
        for c in cmts:
            ws_cmts.append([
                note_id,
                note_title_map.get(note_id, ""),
                c.rank,
                c.content or "",
                c.likes,
            ])

    cmt_widths = [26, 40, 8, 60, 10]
    for i, w in enumerate(cmt_widths, 1):
        ws_cmts.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws_cmts.row_dimensions[1].height = 20
    ws_cmts.freeze_panes = "C2"
    ws_cmts.auto_filter.ref = ws_cmts.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = _scheme_filename(task, "xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.delete("/api/notes/{note_id}")
def delete_note(note_id: str, db: Session = Depends(get_db)):
    note = db.get(Note, note_id)
    if not note:
        raise HTTPException(status_code=404, detail="Note not found")
    db.query(Comment).filter(Comment.note_id == note_id).delete(synchronize_session=False)
    db.delete(note)
    db.commit()
    return {"ok": True}


@router.delete("/api/notes")
def delete_notes_bulk(ids: list[str] = Body(...), db: Session = Depends(get_db)):
    if not ids:
        return {"ok": True, "deleted": 0}
    db.query(Comment).filter(Comment.note_id.in_(ids)).delete(synchronize_session=False)
    deleted = db.query(Note).filter(Note.id.in_(ids)).delete(synchronize_session=False)
    db.commit()
    return {"ok": True, "deleted": deleted}


@router.get("/api/notes/{note_id}")
def get_note(note_id: str, db: Session = Depends(get_db)):
    note = db.get(Note, note_id)
    if not note:
        raise HTTPException(status_code=404, detail="Note not found")
    comments = (
        db.query(Comment)
        .filter(Comment.note_id == note_id)
        .order_by(Comment.rank)
        .all()
    )
    result = NoteDetail.model_validate(note)
    result.comments = [CommentOut.model_validate(c) for c in comments]
    return result
